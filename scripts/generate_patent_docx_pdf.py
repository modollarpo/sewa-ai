"""
Generate Word .docx and PDF from docs/patent/UK_Patent_Application_2_Specification_and_Claims.md
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors

MD_PATH = 'docs/patent/UK_Patent_Application_2_Specification_and_Claims.md'
DOCX_PATH = 'docs/patent/UK_Patent_Application_2_Specification_and_Claims.docx'
PDF_PATH = 'docs/patent/UK_Patent_Application_2_Specification_and_Claims.pdf'

def parse_markdown(path):
    with open(path, 'r', encoding='utf-8') as f:
        text = f.read()
    return text

def build_docx(md_text, out_path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    lines = md_text.split('\n')
    in_claims = False

    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue
        if line_s.startswith('# '):
            p = doc.add_paragraph()
            r = p.add_run(line_s[2:].strip())
            r.bold = True
            r.font.size = Pt(14)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line_s.startswith('## '):
            heading_text = line_s[3:].strip()
            if 'CLAIMS' in heading_text.upper():
                in_claims = True
            doc.add_heading(heading_text, level=1)
        elif line_s.startswith('### '):
            doc.add_heading(line_s[4:].strip(), level=2)
        elif line_s.startswith('---'):
            continue
        elif in_claims and re.match(r'^\d+\.', line_s):
            p = doc.add_paragraph(line_s)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(4)
        elif line_s.startswith('- '):
            p = doc.add_paragraph(line_s[2:], style='List Bullet')
        else:
            p = doc.add_paragraph(line_s)
            p.paragraph_format.space_after = Pt(3)

    doc.save(out_path)
    print(f"Generated DOCX: {out_path}")

def build_pdf(md_text, out_path):
    doc = SimpleDocTemplate(out_path, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle('DocTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=12, leading=15, alignment=1, textColor=colors.HexColor('#050f0c'))
    h1_style = ParagraphStyle('H1', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10.5, leading=13, spaceBefore=8, spaceAfter=3, textColor=colors.HexColor('#050f0c'))
    h2_style = ParagraphStyle('H2', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9.5, leading=12, spaceBefore=6, spaceAfter=2, textColor=colors.HexColor('#374151'))
    body_style = ParagraphStyle('Body', parent=styles['Normal'], fontName='Helvetica', fontSize=8, leading=11, spaceBefore=2, spaceAfter=3, textColor=colors.HexColor('#1f2937'))
    claim_style = ParagraphStyle('Claim', parent=styles['Normal'], fontName='Helvetica', fontSize=7.8, leading=10.5, leftIndent=12, spaceBefore=2, spaceAfter=3, textColor=colors.HexColor('#111827'))

    story = []
    lines = md_text.split('\n')
    in_claims = False

    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue
        if line_s.startswith('# '):
            story.append(Paragraph(line_s[2:].strip(), title_style))
            story.append(Spacer(1, 4))
        elif line_s.startswith('## '):
            heading_text = line_s[3:].strip()
            if 'CLAIMS' in heading_text.upper():
                in_claims = True
                story.append(PageBreak())
            story.append(Paragraph(heading_text, h1_style))
        elif line_s.startswith('### '):
            story.append(Paragraph(line_s[4:].strip(), h2_style))
        elif line_s.startswith('---'):
            story.append(Spacer(1, 4))
        elif in_claims and re.match(r'^\d+\.', line_s):
            clean = line_s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            story.append(Paragraph(clean, claim_style))
        else:
            clean = line_s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            story.append(Paragraph(clean, body_style))

    doc.build(story)
    print(f"Generated PDF: {out_path}")

if __name__ == '__main__':
    md = parse_markdown(MD_PATH)
    build_docx(md, DOCX_PATH)
    build_pdf(md, PDF_PATH)
    print("Patent generation complete.")
