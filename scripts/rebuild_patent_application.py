"""Build institutional-grade UKIPO-submission DOCX and PDF for UK Patent Application 2.

Adds a worked-example paragraph for the Multi-Gate Safety Interception Audit
to Section 7.3 (RMD) and Section 7.4 (Conformal Risk Control), then rebuilds:
  - docs/patent/UK_Patent_Application_2_Specification_and_Claims.md   (master markdown)
  - docs/patent/UK_Patent_Application_2_Specification_and_Claims.docx (spec-only DOCX)
  - docs/patent/UK_Patent_Application_2_Specification_and_Claims.pdf  (spec-only PDF)
  - docs/patent/UK_Patent_Application_2_With_Drawings.docx            (with drawings)
  - docs/patent/UK_Patent_Application_2_With_Drawings.pdf             (with drawings)

FEASIBILITY ON PUBLIC DATA. NOT CLINICAL EVIDENCE.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path
from io import BytesIO

# ── python-docx ──────────────────────────────────────────────────────────────
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── reportlab ─────────────────────────────────────────────────────────────────
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, HRFlowable,
    KeepTogether, PageBreak, Table, TableStyle, Image as RLImage
)
from reportlab.platypus.tableofcontents import TableOfContents

ROOT = Path(__file__).resolve().parents[1]
PATENT_DIR = ROOT / "docs" / "patent"

DRAWING_CAPTIONS = {
    "fig1_architecture.png": "FIG. 1: HIGH-LEVEL SYSTEM ARCHITECTURE & 7-LAYER SAFETY PIPELINE",
    "fig2_conformal_calibration.png": "FIG. 2: SPLIT CONFORMAL RISK CONTROL CALIBRATION & INFERENCE",
    "fig3_adabn_mechanism.png": "FIG. 3: ZERO-LABEL ADAPTIVE BATCH NORMALISATION (AdaBN)",
    "fig4_masked_hirescam.png": "FIG. 4: ANATOMICALLY MASKED HiResCAM ATTRIBUTION",
    "fig5_uncertainty_monotonicity.png": "FIG. 5: EPISTEMIC UNCERTAINTY VARIANCE MONOTONICITY",
    "fig6_stroke_code_dispatch.png": "FIG. 6: STROKE CODE AUTOMATED MULTI-CHANNEL DISPATCH",
    "fig7_sovereign_deployment.png": "FIG. 7: SOVEREIGN DUAL-PATH DEPLOYMENT ARCHITECTURE",
    "fig8_multimodal_expansion.png": "FIG. 8: MULTI-MODAL EMERGENCY RADIOLOGY TRIAGE PIPELINE",
}

FIGURES_DIR = PATENT_DIR / "figures"
SPEC_MD = PATENT_DIR / "UK_Patent_Application_2_Specification_and_Claims.md"

# ─── New worked-example text to insert ────────────────────────────────────────
# Inserted into §7.3 after the existing RMD AUROC sentence.
RMD_WORKED_EXAMPLE = (
    "\n\n*Worked Example — Retrospective Multi-Gate Secondary Safety Audit*: "
    "An exploratory retrospective secondary-gate audit was evaluated on the 6 false-negative studies "
    "(CQ500-CT-385, CQ500-CT-415, CQ500-CT-278, CQ500-CT-426, CQ500-CT-196, CQ500-CT-43) "
    "that scored below lambda = 0.0298 in the single-gate conformal evaluation (conditional FNR = 3.70%, 6/162). "
    "Under single-slice maximum probability aggregation, the six study scores were: "
    "CQ500-CT-385: 0.00962; CQ500-CT-415: 0.00448; CQ500-CT-278: 0.00246; CQ500-CT-426: 0.00627; "
    "CQ500-CT-196: 0.00166; CQ500-CT-43: 0.02181. Under top-3 slice mean probability aggregation, "
    "the corresponding scores were 0.00508, 0.00367, 0.00117, 0.00330, 0.00121, and 0.00959. "
    "Under both aggregation conventions, all 6 positive cases scored strictly below lambda = 0.0298. "
    "To evaluate whether secondary safety gating provides defense-in-depth, an exploratory OOD threshold "
    "tau_OOD = 52.05 was calculated from the 95th percentile of RMD scores across 49 confirmed true-negative "
    "auto-cleared studies within the CQ500 cohort. Under this threshold, 2 of the 6 studies (CQ500-CT-415, RMD = 55.54; "
    "and CQ500-CT-43, RMD = 64.31) exceeded tau_OOD and would have been flagged for human clinical escalation "
    "due to anomalous latent feature representation. Conversely, 4 of the 6 studies (CQ500-CT-385 RMD = 32.43; "
    "CQ500-CT-278 RMD = 41.01; CQ500-CT-426 RMD = 46.58; CQ500-CT-196 RMD = 47.40) yielded RMD scores below "
    "tau_OOD, and none of the six exceeded the 95th-percentile epistemic uncertainty threshold (tau_epi = 3.50e-4). "
    "Because tau_OOD was derived from true-negative cases within the same overall multi-centre cohort, these findings "
    "represent exploratory within-sample analysis rather than prospective or external clinical validation. "
    "Furthermore, an OOD flag reflects feature abnormality rather than a guarantee of clinical triage error prevention. "
    "In this exploratory evaluation, four of six classifier false negatives were not flagged by either evaluated secondary gate. "
    "These results show that the tested gates did not intercept every observed false negative and require further evaluation on independent data. "
    "Consequently, the single-gate classifier FNR of 3.70% (6/162) remains unchanged. While automated routing provides priority worklist "
    "escalation, the platform's architectural design (Claim 1(h)) specifies mandatory consultant radiologist review and sign-off prior to diagnostic release."
)

CONFORMAL_WORKED_EXAMPLE = (
    "- Exploratory multi-gate secondary audit (retrospective, separate analysis): "
    "Of the 6 false-negative studies at lambda = 0.0298, 2 (CQ500-CT-415, CQ500-CT-43) "
    "exhibited RMD scores exceeding an exploratory 95th-percentile true-negative threshold (tau_OOD = 52.05), "
    "triggering human review due to anomalous feature distribution. The remaining 4 studies were not flagged by "
    "either OOD or epistemic uncertainty gates. Because the threshold was derived within-sample, this is an "
    "exploratory analysis; the observed classifier-gate FNR of 3.70% (6/162) remains unchanged. In this exploratory evaluation, four of six classifier false negatives were not flagged by "
    "either evaluated secondary gate. These results show that the tested gates did not intercept every observed false "
    "negative and require further evaluation on independent data. The single-gate classifier FNR of 3.70% (6/162) remains unchanged, "
    "supporting the architectural specification of mandatory consultant radiologist sign-off (Claim 1(h)). Full audit details: validation/reports/v2a1_score_discrepancy_reconciliation.json."
)

def update_markdown(md_path: Path) -> str:
    """Insert or update worked examples into the master markdown. Returns updated content."""
    content = md_path.read_text(encoding="utf-8")

    # Check if worked example is already present
    if "*Worked Example — Retrospective Multi-Gate Secondary Safety Audit*:" in content:
        print("Master markdown already contains worked examples. Keeping synced.")
        return content

    rmd_anchor = (
        "In CQ500 error detection, RMD achieved AUROC 0.5798 vs 0.5256 for standard "
        "Mahalanobis distance."
    )
    if rmd_anchor in content:
        content = content.replace(rmd_anchor, rmd_anchor + RMD_WORKED_EXAMPLE)
        print("§7.3 worked example inserted.")

    conformal_anchor = (
        "- The distinction between theoretical expected risk bounds under exchangeability "
        "and observed empirical sample rates on finite test cohorts is explicitly recognized."
    )
    if conformal_anchor in content:
        content = content.replace(conformal_anchor, conformal_anchor + "\n" + CONFORMAL_WORKED_EXAMPLE)
        print("§7.4 worked example inserted.")

    md_path.write_text(content, encoding="utf-8")
    print(f"Master markdown updated: {md_path.name}")
    return content


# ─────────────────────────────────────────────────────────────────────────────
#  DOCX BUILD
# ─────────────────────────────────────────────────────────────────────────────

def set_doc_defaults(doc: Document) -> None:
    """Apply UKIPO-standard page margins and default font."""
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    # Default Normal style: Times New Roman 12pt, justified
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = Pt(16)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(14)
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(12)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
    else:
        run.font.size = Pt(11)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
    run.font.name = "Times New Roman"


def add_body(doc: Document, text: str, italic: bool = False, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.italic = italic
    run.bold = bold
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = Pt(16)


def add_bullet(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    p.paragraph_format.left_indent = Cm(level * 0.75)
    p.paragraph_format.space_after = Pt(3)


def add_claim_para(doc: Document, number: int, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = Pt(18)
    run_n = p.add_run(f"{number}.  ")
    run_n.font.name = "Times New Roman"
    run_n.font.size = Pt(12)
    run_n.bold = True
    run_t = p.add_run(text)
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(12)


def add_hr(doc: Document) -> None:
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(8)


def add_cover_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("UK PATENT APPLICATION — SPECIFICATION AND CLAIMS")
    r.bold = True; r.font.size = Pt(16); r.font.name = "Times New Roman"
    doc.add_paragraph()

    meta = [
        ("Filing Body", "United Kingdom Intellectual Property Office (UKIPO)"),
        ("Application Type", "Second UK National Patent Application (Patents Act 1977 s.14)"),
        ("Priority Claim", "UK Patent Application No. GB2609706.3, filed 25 April 2026, "
                           "pursuant to s.5(2) Patents Act 1977 / Article 4 Paris Convention"),
        ("Applicant", "STOREGRILL INC LTD (Company No. 14581073)"),
        ("Inventor", "Dolapo Ariyo"),
        ("Address", "18 Hampshire Close, Binley, Coventry, CV3 2FP, United Kingdom"),
        ("Statutory Marking", "FEASIBILITY ON PUBLIC DATA. NOT CLINICAL EVIDENCE. "
                              "PRE-CLINICAL SPECIFICATION (TRL 5 → TRL 6 PLANNED)."),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p.add_run(f"{label}: ")
        r1.bold = True; r1.font.size = Pt(11); r1.font.name = "Times New Roman"
        r2 = p.add_run(value)
        r2.font.size = Pt(11); r2.font.name = "Times New Roman"
        p.paragraph_format.space_after = Pt(3)
    add_hr(doc)


def build_docx_from_markdown(md_content: str, out_path: Path,
                              include_figures: bool = False) -> None:
    """Parse the markdown and produce a properly formatted DOCX."""
    doc = Document()
    set_doc_defaults(doc)
    add_cover_block(doc)

    # Parse markdown line by line with simple state machine
    lines = md_content.splitlines()
    i = 0
    in_claim_section = False
    claim_number = 0

    while i < len(lines):
        line = lines[i]

        # Skip YAML-style header block lines already in cover
        if line.startswith("**Filing Body**") or line.startswith("**Application") or \
           line.startswith("**Priority") or line.startswith("**Applicant") or \
           line.startswith("**Inventor") or line.startswith("**Registered") or \
           line.startswith("**Statutory") or line == "# UK PATENT APPLICATION SPECIFICATION AND CLAIMS":
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ("---", "***", "___"):
            add_hr(doc)
            i += 1
            continue

        # ATX headings
        m = re.match(r"^(#{1,3})\s+(.+)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            # Detect claims section
            if "CLAIMS" in text.upper() or "8. CLAIMS" in text:
                in_claim_section = True
            add_heading(doc, text, level)
            i += 1
            continue

        # Bullet / list
        m_bullet = re.match(r"^(\s*[-*])\s+(.+)$", line)
        if m_bullet:
            indent = 1 + (len(m_bullet.group(1)) - 1) // 2
            add_bullet(doc, m_bullet.group(2).strip(), indent)
            i += 1
            continue

        # Numbered list (claims)
        m_num = re.match(r"^(\d+)\.\s+(.+)$", line)
        if m_num and in_claim_section:
            cn = int(m_num.group(1))
            claim_text = m_num.group(2)
            # Accumulate continuation lines (indented or continuation)
            j = i + 1
            while j < len(lines) and lines[j].startswith("   "):
                claim_text += " " + lines[j].strip()
                j += 1
            add_claim_para(doc, cn, claim_text.strip())
            i = j
            continue

        # Empty line
        if not line.strip():
            i += 1
            continue

        # Inline bold/italic stripping for body paragraphs
        clean = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        clean = re.sub(r"\*(.+?)\*", r"\1", clean)

        # Detect italic-prefixed lines (empirical findings etc.)
        is_italic = line.startswith("*") and line.endswith("*") and not line.startswith("**")
        add_body(doc, clean.strip(), italic=is_italic)
        i += 1

    # Page break before claims
    for para in doc.paragraphs:
        if "8. CLAIMS" in para.text or "WE CLAIM" in para.text:
            para.paragraph_format.page_break_before = True
            break

    doc.save(str(out_path))
    print(f"DOCX saved: {out_path.name}  ({out_path.stat().st_size // 1024} KB)")


# ─────────────────────────────────────────────────────────────────────────────
#  PDF BUILD (ReportLab)
# ─────────────────────────────────────────────────────────────────────────────

PAGE_W, PAGE_H = A4
MARGIN = 2.5 * cm
BODY_W = PAGE_W - 2 * MARGIN

GREY_LIGHT = colors.HexColor("#F5F5F5")
DARK = colors.HexColor("#1A1A1A")
RULE_COLOR = colors.HexColor("#333333")
ACCENT = colors.HexColor("#003087")   # UKIPO navy


def _styles() -> dict:
    base = getSampleStyleSheet()
    s = {}

    s["cover_title"] = ParagraphStyle(
        "cover_title", parent=base["Normal"],
        fontName="Times-Bold", fontSize=16, leading=22,
        alignment=TA_CENTER, spaceAfter=10, textColor=ACCENT
    )
    s["cover_sub"] = ParagraphStyle(
        "cover_sub", parent=base["Normal"],
        fontName="Times-Roman", fontSize=10, leading=14,
        alignment=TA_LEFT, spaceAfter=3
    )
    s["title"] = ParagraphStyle(
        "title", parent=base["Normal"],
        fontName="Times-Bold", fontSize=14, leading=20,
        spaceAfter=8, spaceBefore=16, textColor=ACCENT
    )
    s["h2"] = ParagraphStyle(
        "h2", parent=base["Normal"],
        fontName="Times-Bold", fontSize=12, leading=17,
        spaceAfter=5, spaceBefore=12
    )
    s["h3"] = ParagraphStyle(
        "h3", parent=base["Normal"],
        fontName="Times-Bold", fontSize=11, leading=15,
        spaceAfter=4, spaceBefore=8
    )
    s["body"] = ParagraphStyle(
        "body", parent=base["Normal"],
        fontName="Times-Roman", fontSize=11, leading=17,
        spaceAfter=6, alignment=TA_JUSTIFY
    )
    s["body_italic"] = ParagraphStyle(
        "body_italic", parent=base["Normal"],
        fontName="Times-Italic", fontSize=11, leading=17,
        spaceAfter=6, alignment=TA_JUSTIFY
    )
    s["bullet"] = ParagraphStyle(
        "bullet", parent=base["Normal"],
        fontName="Times-Roman", fontSize=10.5, leading=15,
        spaceAfter=3, leftIndent=16, bulletIndent=0, firstLineIndent=0,
        alignment=TA_JUSTIFY
    )
    s["claim"] = ParagraphStyle(
        "claim", parent=base["Normal"],
        fontName="Times-Roman", fontSize=11, leading=18,
        spaceAfter=10, leftIndent=14, firstLineIndent=-14,
        alignment=TA_JUSTIFY
    )
    s["claim_num"] = ParagraphStyle(
        "claim_num", parent=base["Normal"],
        fontName="Times-Bold", fontSize=11, leading=18,
        spaceAfter=0
    )
    s["marking"] = ParagraphStyle(
        "marking", parent=base["Normal"],
        fontName="Times-Italic", fontSize=8, leading=11,
        textColor=colors.HexColor("#666666"), alignment=TA_CENTER
    )
    return s


def _header_footer(canvas, doc, title: str) -> None:
    canvas.saveState()
    canvas.setFont("Times-Roman", 8)
    canvas.setFillColor(colors.HexColor("#666666"))
    # Header
    canvas.drawString(MARGIN, PAGE_H - MARGIN * 0.6,
                      "UK PATENT APPLICATION — SPECIFICATION AND CLAIMS")
    canvas.drawRightString(PAGE_W - MARGIN, PAGE_H - MARGIN * 0.6, title[:60])
    canvas.setStrokeColor(RULE_COLOR)
    canvas.setLineWidth(0.4)
    canvas.line(MARGIN, PAGE_H - MARGIN * 0.7, PAGE_W - MARGIN, PAGE_H - MARGIN * 0.7)
    # Footer
    canvas.line(MARGIN, MARGIN * 0.85, PAGE_W - MARGIN, MARGIN * 0.85)
    canvas.drawString(MARGIN, MARGIN * 0.55, "FEASIBILITY ON PUBLIC DATA. NOT CLINICAL EVIDENCE.")
    canvas.drawCentredString(PAGE_W / 2, MARGIN * 0.55, f"Page {doc.page}")
    canvas.drawRightString(PAGE_W - MARGIN, MARGIN * 0.55, "STOREGRILL INC LTD — CONFIDENTIAL")
    canvas.restoreState()


def clean_for_xml(text: str) -> str:
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    return text


def build_pdf_from_markdown(md_content: str, out_path: Path,
                             include_figures: bool = False,
                             doc_title: str = "UK Patent Application 2") -> None:
    s = _styles()
    story = []

    def hr():
        story.append(HRFlowable(width="100%", thickness=0.5,
                                 color=RULE_COLOR, spaceAfter=6, spaceBefore=6))

    def cover() -> None:
        story.append(Spacer(1, 1.5 * cm))
        story.append(Paragraph("UNITED KINGDOM INTELLECTUAL PROPERTY OFFICE", s["marking"]))
        story.append(Spacer(1, 0.3 * cm))
        story.append(Paragraph(
            "UK PATENT APPLICATION — SPECIFICATION AND CLAIMS", s["cover_title"]
        ))
        story.append(Spacer(1, 0.5 * cm))
        hr()
        meta = [
            ("<b>Title:</b>",
             "Clinical Artificial Intelligence Platform and Method for Distribution-Free "
             "Statistical Risk-Controlled Radiological Triage, Zero-Label Domain Adaptation, "
             "and Anatomically-Bounded Attribution"),
            ("<b>Filing Body:</b>", "United Kingdom Intellectual Property Office (UKIPO)"),
            ("<b>Application Type:</b>",
             "Second UK National Patent Application (Patents Act 1977 s.14)"),
            ("<b>Priority Claim:</b>",
             "UK Patent Application No. GB2609706.3, filed 25 April 2026, "
             "pursuant to s.5(2) Patents Act 1977 / Article 4 Paris Convention"),
            ("<b>Applicant:</b>", "STOREGRILL INC LTD (Company No. 14581073)"),
            ("<b>Inventor:</b>", "Dolapo Ariyo"),
            ("<b>Address:</b>", "18 Hampshire Close, Binley, Coventry, CV3 2FP, United Kingdom"),
            ("<b>Statutory Marking:</b>",
             "FEASIBILITY ON PUBLIC DATA. NOT CLINICAL EVIDENCE. "
             "PRE-CLINICAL SPECIFICATION (TRL 5 → TRL 6 PLANNED)."),
        ]
        data = [[Paragraph(k, s["cover_sub"]), Paragraph(v, s["cover_sub"])] for k, v in meta]
        t = Table(data, colWidths=[3.8 * cm, BODY_W - 3.8 * cm])
        t.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t)
        hr()
        story.append(PageBreak())

    cover()

    # Parse markdown
    lines = md_content.splitlines()
    i = 0
    in_claim_section = False
    skip_cover_lines = {
        "# UK PATENT APPLICATION SPECIFICATION AND CLAIMS",
    }

    while i < len(lines):
        line = lines[i]

        # Skip cover lines already rendered
        if any(line.startswith(p) for p in (
            "**Filing Body**", "**Application", "**Priority", "**Applicant",
            "**Inventor", "**Registered", "**Statutory"
        )) or line in skip_cover_lines:
            i += 1
            continue

        # HR
        if line.strip() in ("---", "***", "___"):
            hr()
            i += 1
            continue

        # Headings
        m = re.match(r"^(#{1,3})\s+(.+)$", line)
        if m:
            level = len(m.group(1))
            text = clean_for_xml(m.group(2).strip())
            if "CLAIMS" in text.upper() or "8. CLAIMS" in text:
                in_claim_section = True
                story.append(PageBreak())
            sty = s["title"] if level == 1 else s["h2"] if level == 2 else s["h3"]
            story.append(Paragraph(text, sty))
            i += 1
            continue

        # Bullets
        m_bullet = re.match(r"^(\s*[-*])\s+(.+)$", line)
        if m_bullet:
            txt = clean_for_xml(m_bullet.group(2).strip())
            story.append(Paragraph(f"• {txt}", s["bullet"]))
            i += 1
            continue

        # Numbered list items in claims section
        m_num = re.match(r"^(\d+)\.\s+(.+)$", line)
        if m_num and in_claim_section:
            cn = m_num.group(1)
            claim_text = m_num.group(2)
            j = i + 1
            while j < len(lines) and lines[j].startswith("   "):
                claim_text += " " + lines[j].strip()
                j += 1
            claim_text = clean_for_xml(claim_text.strip())
            story.append(Paragraph(f"<b>{cn}.</b>  {claim_text}", s["claim"]))
            i = j
            continue

        # Empty line
        if not line.strip():
            story.append(Spacer(1, 4))
            i += 1
            continue

        # Body paragraph
        txt = line.strip()
        is_italic = txt.startswith("*") and txt.endswith("*") and not txt.startswith("**")
        txt = clean_for_xml(txt)
        sty = s["body_italic"] if is_italic else s["body"]
        try:
            story.append(Paragraph(txt, sty))
        except Exception:
            story.append(Paragraph(txt.encode("ascii", "replace").decode(), sty))
        i += 1

    # If figures requested, append them after spec
    if include_figures and FIGURES_DIR.exists():
        fig_files = sorted(FIGURES_DIR.glob("*.png"))
        for idx, fig_path in enumerate(fig_files):
            story.append(PageBreak())
            cap_text = DRAWING_CAPTIONS.get(fig_path.name, fig_path.stem.replace("_", " ").upper())
            story.append(Paragraph(f"PATENT DRAWINGS SHEET {idx+1}/{len(fig_files)} — {cap_text}", s["h3"]))
            story.append(Spacer(1, 0.3 * cm))
            # Image aspect ratio 2036 x 2601 (height/width = 1.277)
            # 13.5 cm width * 1.277 = 17.24 cm height fits page without clipping
            img = RLImage(str(fig_path), width=13.5 * cm, height=17.24 * cm)
            story.append(img)

    def first_page(canvas, doc):
        _header_footer(canvas, doc, doc_title)

    def later_pages(canvas, doc):
        _header_footer(canvas, doc, doc_title)

    pdf = SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=MARGIN * 1.5, bottomMargin=MARGIN * 1.5,
        title=doc_title,
        author="Dolapo Ariyo / STOREGRILL INC LTD",
        subject="UK Patent Application — Clinical AI Triage Platform",
        creator="SEWA XAI Patent Build System v2",
    )
    pdf.build(story, onFirstPage=first_page, onLaterPages=later_pages)
    print(f"PDF saved:  {out_path.name}  ({out_path.stat().st_size // 1024} KB)")


# ─────────────────────────────────────────────────────────────────────────────
#  MAIN
# ─────────────────────────────────────────────────────────────────────────────

def main() -> None:
    print("=" * 72)
    print("SEWA XAI Patent Application 2 — Institutional-Grade Rebuild")
    print("FEASIBILITY ON PUBLIC DATA. NOT CLINICAL EVIDENCE.")
    print("=" * 72)

    # 1. Update master markdown
    print("\n[1/5] Inserting worked examples into master markdown...")
    md_content = update_markdown(SPEC_MD)

    # 2. Spec-only DOCX
    print("\n[2/5] Building spec-only DOCX...")
    spec_docx = PATENT_DIR / "UK_Patent_Application_2_Specification_and_Claims.docx"
    build_docx_from_markdown(md_content, spec_docx, include_figures=False)

    # 3. Spec-only PDF
    print("\n[3/5] Building spec-only PDF...")
    spec_pdf = PATENT_DIR / "UK_Patent_Application_2_Specification_and_Claims.pdf"
    build_pdf_from_markdown(
        md_content, spec_pdf, include_figures=False,
        doc_title="UK Patent Application 2 — Specification & Claims"
    )

    # 4. With-drawings DOCX (spec + figure embed references)
    print("\n[4/5] Building with-drawings DOCX...")
    draw_docx = PATENT_DIR / "UK_Patent_Application_2_With_Drawings.docx"
    
    # Check if draw_docx is locked by another process
    target_docx = draw_docx
    try:
        with open(draw_docx, 'a+b') as test_lock:
            pass
    except (PermissionError, OSError):
        print(f"Notice: {draw_docx.name} is currently open by the user. Generating in-place as soon as available or to fallback.")
        target_docx = draw_docx

    # Build directly in a temporary Document object in memory
    doc_with = Document()
    set_doc_defaults(doc_with)
    add_cover_block(doc_with)

    # Parse and add text
    lines = md_content.splitlines()
    i = 0
    in_claim_section = False
    while i < len(lines):
        line = lines[i]
        if line.startswith("**Filing Body**") or line.startswith("**Application") or \
           line.startswith("**Priority") or line.startswith("**Applicant") or \
           line.startswith("**Inventor") or line.startswith("**Registered") or \
           line.startswith("**Statutory") or line == "# UK PATENT APPLICATION SPECIFICATION AND CLAIMS":
            i += 1
            continue
        if line.strip() in ("---", "***", "___"):
            add_hr(doc_with)
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.+)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if "CLAIMS" in text.upper() or "8. CLAIMS" in text:
                in_claim_section = True
            add_heading(doc_with, text, level)
            i += 1
            continue
        m_bullet = re.match(r"^(\s*[-*])\s+(.+)$", line)
        if m_bullet:
            indent = 1 + (len(m_bullet.group(1)) - 1) // 2
            add_bullet(doc_with, m_bullet.group(2).strip(), indent)
            i += 1
            continue
        m_num = re.match(r"^(\d+)\.\s+(.+)$", line)
        if m_num and in_claim_section:
            cn = int(m_num.group(1))
            claim_text = m_num.group(2)
            j = i + 1
            while j < len(lines) and lines[j].startswith("   "):
                claim_text += " " + lines[j].strip()
                j += 1
            add_claim_para(doc_with, cn, claim_text.strip())
            i = j
            continue
        if not line.strip():
            i += 1
            continue
        clean = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        clean = re.sub(r"\*(.+?)\*", r"\1", clean)
        is_italic = line.startswith("*") and line.endswith("*") and not line.startswith("**")
        add_body(doc_with, clean.strip(), italic=is_italic)
        i += 1

    for para in doc_with.paragraphs:
        if "8. CLAIMS" in para.text or "WE CLAIM" in para.text:
            para.paragraph_format.page_break_before = True
            break

    # Append formal drawings
    if FIGURES_DIR.exists():
        fig_files = sorted(FIGURES_DIR.glob("*.png"))
        for idx, fig_path in enumerate(fig_files):
            doc_with.add_page_break()

            cap_text = DRAWING_CAPTIONS.get(fig_path.name, fig_path.stem.replace("_", " ").upper())

            # Header paragraph with single spacing
            p_hdr = doc_with.add_paragraph()
            p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_hdr.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p_hdr.paragraph_format.line_spacing = 1.0
            p_hdr.paragraph_format.space_before = Pt(0)
            p_hdr.paragraph_format.space_after = Pt(6)
            r_hdr = p_hdr.add_run(f"PATENT DRAWINGS SHEET {idx+1}/{len(fig_files)}\n{cap_text}")
            r_hdr.bold = True
            r_hdr.font.size = Pt(11)
            r_hdr.font.name = "Times New Roman"

            # Image paragraph with CRITICAL SINGLE LINE SPACING to prevent Word 16pt clipping
            p_img = doc_with.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p_img.paragraph_format.line_spacing = 1.0
            p_img.paragraph_format.space_before = Pt(0)
            p_img.paragraph_format.space_after = Pt(0)

            run_img = p_img.add_run()
            # 5.4 inches width (height = 5.4 * 2601 / 2036 = 6.89 inches)
            run_img.add_picture(str(fig_path), width=Inches(5.4))

    try:
        doc_with.save(str(draw_docx))
        print(f"With-drawings DOCX saved: {draw_docx.name}  ({draw_docx.stat().st_size // 1024} KB)")
    except (PermissionError, OSError):
        alt_docx = draw_docx.with_name("UK_Patent_Application_2_With_Drawings_updated.docx")
        doc_with.save(str(alt_docx))
        print(f"With-drawings DOCX saved to {alt_docx.name} (original locked in Word) ({alt_docx.stat().st_size // 1024} KB)")

    # 5. With-drawings PDF
    print("\n[5/5] Building with-drawings PDF...")
    draw_pdf = PATENT_DIR / "UK_Patent_Application_2_With_Drawings.pdf"
    try:
        build_pdf_from_markdown(
            md_content, draw_pdf, include_figures=True,
            doc_title="UK Patent Application 2 — Specification, Claims & Drawings"
        )
    except (PermissionError, OSError):
        alt_pdf = draw_pdf.with_name("UK_Patent_Application_2_With_Drawings_updated.pdf")
        build_pdf_from_markdown(
            md_content, alt_pdf, include_figures=True,
            doc_title="UK Patent Application 2 — Specification, Claims & Drawings"
        )

    print("\n" + "=" * 72)
    print("ALL OUTPUTS:")
    for p in [SPEC_MD, spec_docx, spec_pdf, draw_docx, draw_pdf]:
        size = p.stat().st_size
        print(f"  {p.name:<60} {size/1024:>8.1f} KB")
    print("=" * 72)
    print("FEASIBILITY ON PUBLIC DATA. NOT CLINICAL EVIDENCE.")


if __name__ == "__main__":
    main()
