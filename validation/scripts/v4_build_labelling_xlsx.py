"""v4 PART A: build the unlabelled .xlsx labelling sheets + private key + hashes.

FEASIBILITY ON PUBLIC DATA. NOT CLINICAL EVIDENCE.

Creates (no labels exist yet; the private key is kept out of every shared sheet):
  validation/labelling/sewa_labelling_sheet.xlsx       - full 300 rows (variant-blind)
  validation/labelling/dolapo_subset_sheet.xlsx        - the 50 second-labeller rows
  validation/labelling/clinician_review_sheet.xlsx     - extractor sentences (no key)
  validation/labelling/private/extractor_sentences_key.csv - PRIVATE key (proposed labels)
  docs/evidence/label_hashes.txt                       - SHA-256 of the above

Each sheet has a dropdown (data validation) on its label column, a frozen header row,
wrapped report text, and an "Instructions" first sheet.

Usage: python -m validation.scripts.v4_build_labelling_xlsx
"""
from __future__ import annotations

import csv
import hashlib
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MARK = "FEASIBILITY ON PUBLIC DATA. NOT CLINICAL EVIDENCE."
MERGED = Path("validation/reports/v4_g1_merged_labelling_sheet.csv")
CLIN_CSV = Path("validation/reports/v4_extractor_clinician_review_clinician.csv")
KEY_SRC = Path("validation/reports/v4_extractor_clinician_review.csv")

OUT_DIR = Path("validation/labelling")
PRIV_DIR = OUT_DIR / "private"
SEWA = OUT_DIR / "sewa_labelling_sheet.xlsx"
DOLAPO = OUT_DIR / "dolapo_subset_sheet.xlsx"
CLIN_XLSX = OUT_DIR / "clinician_review_sheet.xlsx"
KEY_PRIV = PRIV_DIR / "extractor_sentences_key.csv"
HASHES = Path("docs/evidence/label_hashes.txt")

LABEL_VALUES = ["consistent", "contradicts", "ambiguous"]
CLIN_VALUES = ["present", "absent", "uncertain", "historical", "resolved"]

RULE = ("If a report both asserts and denies the finding, label it ambiguous and add a note. "
        "label = consistent | contradicts | ambiguous. 'contradicts' means the report asserts "
        "the OPPOSITE of the triage decision shown in the 'decision' column.")
COLUMNS = ["new_id", "decision", "probability", "urgency", "report_text", "label", "notes",
           "instructions"]


def _instructions_sheet(wb: Workbook, title_lines: list[str]) -> None:
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 120
    for i, line in enumerate(title_lines, start=1):
        c = ws.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def _data_sheet(wb: Workbook, cols: list[str], rows: list[list], label_col: str,
                label_values: list[str], wrap_col: str = "report_text") -> None:
    ws = wb.create_sheet("Labels")
    ws.append(cols)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    # wrap the report text column
    if wrap_col in cols:
        ci = cols.index(wrap_col) + 1
        ws.column_dimensions[get_column_letter(ci)].width = 80
        for row in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    # dropdown on the label column
    li = cols.index(label_col) + 1
    letter = get_column_letter(li)
    dv = DataValidation(type="list", formula1='"' + ",".join(label_values) + '"',
                        allow_blank=True, showDropDown=False)
    dv.error = "Choose: " + ", ".join(label_values)
    dv.errorTitle = "Invalid label"
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{max(len(rows) + 1, 2)}")


def build_sewa(rows: list[dict]) -> None:
    wb = Workbook()
    _instructions_sheet(wb, [MARK, "", "SEWA - report-negation labelling", "", RULE,
                             "", "Columns: new_id, decision, probability, urgency, "
                             "report_text, label, notes, instructions."])
    data = [[r["new_id"], r["decision"], r["probability"], r["urgency"], r["report_text"],
             "", "", r["instructions"]] for r in rows]
    _data_sheet(wb, COLUMNS, data, "label", LABEL_VALUES)
    wb.save(SEWA)


def build_dolapo(rows: list[dict]) -> None:
    wb = Workbook()
    _instructions_sheet(wb, [MARK, "", "Second labeller (subset) - report-negation labelling",
                             "", RULE, "", "These are the 50 rows flagged for a second "
                             "labeller. Apply your own judgement; no other labels are shown."])
    data = [[r["new_id"], r["decision"], r["probability"], r["urgency"], r["report_text"],
             "", "", r["instructions"]] for r in rows]
    _data_sheet(wb, COLUMNS, data, "label", LABEL_VALUES)
    wb.save(DOLAPO)


def build_clinician() -> int:
    with CLIN_CSV.open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    cols = ["id", "sentence", "clinician_label", "comment"]
    wb = Workbook()
    _instructions_sheet(wb, [MARK, "", "Clinician review - haemorrhage polarity of one sentence",
                             "", "For each sentence choose the sentence's haemorrhage status:",
                             "present | absent | uncertain | historical | resolved",
                             "(the exact extractor category names).", "",
                             "This sheet contains NO proposed labels."])
    data = [[r["id"], r["sentence"], "", r.get("comment", "")] for r in rows]
    _data_sheet(wb, cols, data, "clinician_label", CLIN_VALUES, wrap_col="sentence")
    wb.save(CLIN_XLSX)
    return len(rows)


def copy_private_key() -> None:
    PRIV_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(KEY_SRC, KEY_PRIV)


def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def write_hashes(files: list[Path]) -> None:
    lines = [f"# {MARK}", "# SHA-256 of unlabelled labelling artefacts and the private key", ""]
    for p in files:
        lines.append(f"{sha256(p) if p.exists() else 'MISSING':<64}  {p.as_posix()}")
    HASHES.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with MERGED.open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    subset = [r for r in rows if (r.get("second_labeller_subset", "").strip().upper() == "TRUE")]
    print(f"{MARK}\nmerged rows={len(rows)}  second-labeller subset={len(subset)}")

    build_sewa(rows)
    build_dolapo(subset)
    n_clin = build_clinician()
    copy_private_key()
    print(f"clinician rows={n_clin}")

    files = [CLIN_CSV, Path("validation/reports/v4_g1_blinded_labelling_sheet.csv"),
             Path("validation/reports/v4_g1b_blinded_labelling_sheet.csv"), MERGED,
             SEWA, DOLAPO, CLIN_XLSX, KEY_PRIV]
    write_hashes(files)

    for p in (SEWA, DOLAPO, CLIN_XLSX, KEY_PRIV):
        print(f"wrote {p}")
    print(f"wrote {HASHES}")
    for ln in HASHES.read_text(encoding="utf-8").splitlines()[2:]:
        print("  ", ln)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
